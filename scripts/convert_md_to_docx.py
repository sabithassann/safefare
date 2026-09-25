import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MD_PATH = r"c:\Github\safefare\doc\FINAL_THESIS_REPORT_SAFEFARE.md"
DOCX_PATH = r"c:\Github\safefare\doc\FINAL_THESIS_REPORT_SAFEFARE.docx"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_inline_runs(paragraph, text, base_font_size=11, base_color=RGBColor(30, 41, 59)):
    pattern = re.compile(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|[^\*`\[]+)')
    tokens = pattern.findall(text)
    
    for token in tokens:
        if token.startswith('`') and token.endswith('`') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('**') and token.endswith('**') and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.bold = True
            run.font.color.rgb = base_color
        elif token.startswith('*') and token.endswith('*') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.italic = True
            run.font.color.rgb = base_color
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                link_text = m.group(1)
                run = paragraph.add_run(link_text)
                run.font.name = 'Calibri'
                run.font.size = Pt(base_font_size)
                run.underline = True
                run.font.color.rgb = RGBColor(37, 99, 235)
            else:
                run = paragraph.add_run(token)
                run.font.name = 'Calibri'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = base_color
        else:
            run = paragraph.add_run(token)
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = base_color

def convert_markdown_to_docx(md_path, docx_path):
    doc = Document()

    # Set 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_rows = []

    i = 0
    n = len(lines)

    while i < n:
        raw_line = lines[i]
        line = raw_line.rstrip('\r\n')
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                code_text = '\n'.join(code_block_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, 'F1F5F9')
                set_cell_margins(cell, 120, 120, 150, 150)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                code_block_lines = []
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            else:
                in_code_block = True
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            is_separator = all(c in '|-: ' for c in stripped)
            if not is_separator:
                cols = [c.strip() for c in stripped.strip('|').split('|')]
                table_rows.append(cols)
            in_table = True
            i += 1
            continue
        elif in_table:
            in_table = False
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True
                
                for r_idx, row_data in enumerate(table_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = tbl.cell(r_idx, c_idx)
                        set_cell_margins(cell, 100, 100, 120, 120)
                        val = row_data[c_idx] if c_idx < len(row_data) else ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        
                        if is_header:
                            set_cell_background(cell, '1E293B')
                            add_inline_runs(p, val, base_font_size=10, base_color=RGBColor(255, 255, 255))
                            if p.runs:
                                p.runs[0].bold = True
                        else:
                            bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
                            set_cell_background(cell, bg)
                            add_inline_runs(p, val, base_font_size=9.5, base_color=RGBColor(30, 41, 59))
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            table_rows = []

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal Rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '8')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'CBD5E1')
            p_border.append(bottom)
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            if text.startswith('CHAPTER'):
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
        elif stripped.startswith('## '):
            text = stripped[3:].strip()
            if text.startswith('CHAPTER'):
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(14.5)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue
        elif stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Blockquote / Alert
        if stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, quote_text, base_font_size=10.5, base_color=RGBColor(71, 85, 105))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Lists (bulleted or numbered)
        if stripped.startswith(('- ', '* ')):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, bullet_text, base_font_size=11)
            i += 1
            continue
        
        # Numbered list
        m_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m_num:
            item_text = m_num.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, item_text, base_font_size=11)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        
        # Center title page text if applicable
        if any(keyword in stripped for keyword in ["A Thesis Report Submitted", "DECLARATION OF AUTHORSHIP", "CERTIFICATE OF APPROVAL"]):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        add_inline_runs(p, stripped, base_font_size=11)
        i += 1

    # Ensure in_table is closed if ended with table
    if in_table and table_rows:
        num_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        for r_idx, row_data in enumerate(table_rows):
            is_header = (r_idx == 0)
            for c_idx in range(num_cols):
                cell = tbl.cell(r_idx, c_idx)
                set_cell_margins(cell, 100, 100, 120, 120)
                val = row_data[c_idx] if c_idx < len(row_data) else ""
                p = cell.paragraphs[0]
                if is_header:
                    set_cell_background(cell, '1E293B')
                    add_inline_runs(p, val, base_font_size=10, base_color=RGBColor(255, 255, 255))
                    if p.runs: p.runs[0].bold = True
                else:
                    bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
                    set_cell_background(cell, bg)
                    add_inline_runs(p, val, base_font_size=9.5, base_color=RGBColor(30, 41, 59))

    doc.save(docx_path)
    print(f"Successfully converted {md_path} to {docx_path}")

if __name__ == '__main__':
    convert_markdown_to_docx(MD_PATH, DOCX_PATH)
